"""Spreadsheet export builder (GET + POST).

Split verbatim out of app.py (v9.8.1 modularisation).
Every statement is byte-for-byte identical to app.py; only the route
decorators changed from @app.<verb> to @router.<verb>.
Audit map of source line spans: MIGRATED_LINES.txt
"""
from ._core import get_db
from .fundamental import get_fundamental
from fastapi import APIRouter
from fastapi import Query
from fastapi.responses import Response
import pandas as pd


router = APIRouter()


def _disposition(name: str) -> str:
    """Content-Dispositionِ امن — نامِ فارسی نباید هدرِ لاتین-۱ را بشکند.

    خطایِ قبلی روی هر نمادِ فارسی:
      'latin-1' codec can't encode characters … (کلِ خروجی بنیادی می‌مرد)
    راه‌حل: `filename=` با معادلِ ASCII و `filename*=` با RFC 5987 (UTF-8).
    """
    import unicodedata
    from urllib.parse import quote
    ascii_name = (unicodedata.normalize("NFKD", name)
                  .encode("ascii", "ignore").decode("ascii"))
    safe = "".join(c if c.isalnum() or c in "._-" else "_" for c in ascii_name)
    safe = (safe or "export")[:120]
    return 'attachment; filename="%s"; filename*=UTF-8\'\'%s' % (safe, quote(name))


def _export_build(fmt: str, title: str, rows_data):
    """ساخت واقعی فایل خروجی از rows آماده (پس از اعمال فیلترها) + سازگاری GET قدیمی"""
    import io
    import datetime as _dt
    rows = rows_data
    ts = _dt.datetime.now().strftime("%Y%m%d_%H%M")
    name = f"borsagent_{title}_{ts}"
    if fmt in ("docx", "pdf") and len(rows) > 201:
        header = rows[0]
        rows = header + [["..."] * len(header)] + rows[1:201]

    if fmt == "xlsx":
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "export"
        for r in rows:
            ws.append(["" if c is None else c for c in r])
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 22
        buf = io.BytesIO()
        wb.save(buf)
        return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": _disposition(name + ".xlsx")})
    elif fmt == "docx":
        import docx
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        d = docx.Document()
        d.add_heading(f"BorsTerminal - {title}", 0)
        if rows:
            ncols = max(len(r) for r in rows)
            tblEl = OxmlElement('w:tbl')
            grid = OxmlElement('w:tblGrid')
            for _ in range(ncols):
                grid.append(OxmlElement('w:gridCol'))
            tblEl.append(grid)
            styleEl = OxmlElement('w:tblPr')
            b = OxmlElement('w:tblBorders'); b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4'); b.set(qn('w:color'), '94a3b8')
            styleEl.append(b)
            tblEl.append(styleEl)
            for r in rows:
                tr = OxmlElement('w:tr')
                for c in r:
                    tc = OxmlElement('w:tc')
                    tcp = OxmlElement('w:tcPr')
                    w_el = OxmlElement('w:tcW'); w_el.set(qn('w:w'), '2200'); tcp.append(w_el)
                    tc.append(tcp)
                    p = OxmlElement('w:p')
                    rEl = OxmlElement('w:r')
                    txt = OxmlElement('w:t'); txt.text = str("" if c is None else c)
                    rEl.append(txt); p.append(rEl); tc.append(p)
                    tr.append(tc)
                tblEl.append(tr)
            d._element.body.append(tblEl)
        buf = io.BytesIO()
        d.save(buf)
        return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        headers={"Content-Disposition": _disposition(name + ".docx")})
    elif fmt == "pdf":
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=18, rightMargin=18, topMargin=18, bottomMargin=18)
        styles = getSampleStyleSheet()
        story = [Paragraph(f"<b>BorsTerminal - {title}</b>", styles["Title"]), Spacer(1, 8)]
        if rows:
            data = [[str("" if c is None else c) for c in r] for r in rows]
            tbl = Table(data, repeatRows=1)
            tbl.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0ea5e9")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#94a3b8")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
            ]))
            story.append(tbl)
        doc.build(story)
        return Response(buf.getvalue(), media_type="application/pdf",
                        headers={"Content-Disposition": _disposition(name + ".pdf")})
    else:
        import csv as _csv
        buf = io.StringIO()
        w = _csv.writer(buf)
        w.writerows(rows)
        return Response("\ufeff" + buf.getvalue(), media_type="text/csv; charset=utf-8",
                        headers={"Content-Disposition": _disposition(name + ".csv")})

@router.post("/api/export")
def export_data_post(payload: dict = None):
    """دانلود خروجی با دادهٔ بستهبندیشده از UI (فیلترهای فعال اعمال شده)
    body: {fmt, title, rows: [[...], ...]}"""
    if payload is None:
        return {"status": "error", "message": "empty body"}
    return _export_build(payload.get("fmt", "xlsx"), payload.get("title", "export"), payload.get("rows") or [])

@router.get("/api/export")
def export_data(fmt: str = Query("csv"), view: str = Query("market"), symbol: str = Query(None)):
    """دانلود خروجی قدیمی (بدون فیلتر): market|fundamental — برای backward compat"""
    conn = get_db()
    try:
        rows = []
        if view == "fundamental" and symbol:
            j = get_fundamental(symbol)
            # سطرِ جدولِ شاخص ۲ (سلول‌ها + وضعیتِ ناقص) — همان منبعِ UI
            _pr = ((j.get("details") or {}).get("۲") or {}).get("period_row") or {}
            if j.get("metrics"):
                m = j["metrics"]
                rows.append(["شاخص", "مقدار"])
                rows.append(["ارزش بازار (ریال)", m.get("mcap")])
                rows.append(["فروش (میلیون ریال)", m.get("revenue")])
                rows.append(["حاشیه ناخالص ٪", m.get("gross_margin")])
                rows.append(["ROE ٪", m.get("roe")])
                rows.append(["نسبت P/S", m.get("ps")])
                rows.append(["تعداد صورت‌های مالی", j.get("fs_count")])
                rows.append(["مرجع", j.get("ref_symbol") or "-"])
            for ins in j.get("insights", []):
                lbl = f"شاخص {ins.get('step','')}".strip()
                txt = ins.get("text")
                if str(ins.get("step", "")) == "۲" and _pr.get("cells"):
                    # قاعدهٔ خروجی جدول: سطر حذف نمیشود — مقادیرِ موجود با «-» برایِ
                    # دورهٔ ناموجود، و ذکرِ تعدادِ دوره‌ها در عنوانِ سطر.
                    if _pr.get("red"):
                        lbl = _pr.get("label") or lbl
                    txt = "%s · EPS: %s" % (txt, " | ".join(_pr["cells"]))
                rows.append([lbl, ins.get("title"), txt])
        else:
            df = pd.read_sql_query("""
                SELECT i.l_val18 AS symbol, i.l_val30 AS name,
                       COALESCE(NULLIF(i.sector_name, ''), 'سایر') AS sector,
                       m.p_closing AS closing, m.p_last AS last,
                       m.price_change AS change, m.q_tot_tran AS volume,
                       m.q_tot_cap AS value, m.pe, m.eps, m.z_tot_tran AS trades
                FROM market_watch m LEFT JOIN instruments i ON i.ins_code = m.ins_code
                WHERE i.l_val18 IS NOT NULL
                ORDER BY i.l_val18
            """, conn)
            rows = [df.columns.tolist()] + df.values.tolist()
        return _export_build(fmt, f"{view}_{symbol or ''}", rows)
    except Exception as e:
        return {"status": "error", "message": str(e)}
    finally:
        conn.close()
