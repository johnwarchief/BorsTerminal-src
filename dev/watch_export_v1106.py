#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""dev/watch_export_v1106.py — گاردِ خروجیِ واقعیِ دیدهبان (Excel و CSV).

چرا این گارد: مأموریت می‌گوید «دکمه نباید فقط ظاهر شود یا به لینک ساختگی اشاره
کند» و «فایل دانلودشده را واقعاً در تست باز کن». پس اینجا فایل **ساخته و باز
می‌شود** — نه ازِ رابط، ازِ همان کارخانه‌ای که رابط صدا می‌زند
(`api/_export.py::_export_build`)، با همان بدنه‌ای که `POST /api/export`
می‌گیرد. بی‌شبکه، بی‌نوشتن رویِ `market.db`، بی‌کشتنِ هیچ پروسه‌ای.

چه قفلی می‌زند:
  ۱) عدد درِ Excel سلولِ عددی است (نه «۱٬۲۳۴»ِ رشته‌ای) — وگرنه sort/جمع درِ
     Excel می‌میرد و کاربر آن را «خروجیِ خراب» می‌خواند.
  ۲) `None` تهی می‌ماند؛ صفر نمی‌شود (صفرِ جعلی یعنی «حقیقی هیچی نخرید»).
  ۳) سرستونِ فارسی و واحدِ روشن درِ فایل می‌ماند (۳۰ کاراکتر و ریال/تومان
     مبهم نمی‌شوند).
  ۴) ردپایِ منشأ/تازگی درِ صفحۀ دومِ Excel و درِ انتهای CSV است، و ردیف‌هایِ
     داده را دست نمی‌دهد.
  ۵) CSV با BOMِ UTF-8 و جداکنندۀ پایدار بیرون می‌آید (Excelِ فارسی بی‌BOM
     یونیکد را می‌سوزاند).
  ۶) نامِ فایل RFC 5987 است: `filename=` برایِ ASCII و `filename*=` برایِ همان
     برچسبِ فارسی — همان نقصی که کلِ خروجیِ بنیادی را می‌کُشت.
  ۷) بی‌`meta` رفتارِ پیشین دست‌نخورده می‌ماند (یکِ صفحۀ «export»).
  ۸) بدنه‌هایِ نامعتبر رد می‌شوند: `rows` غیرفهرستی ⇒ ۴۲۲، بیش از سقف ⇒ ۴۱۳.
  ۹) هیچ رازِ نشستی (token/cookie/…) درِ بدنه یا نامِ فایل راه نمی‌یابد.
"""
from __future__ import annotations

import csv
import io
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

import openpyxl                                                # noqa: E402
from api import _export                                        # noqa: E402

PASS = 0
FAIL = 0


def ck(cond, label, detail=None):
    global PASS, FAIL
    if cond:
        PASS += 1
        print("  ok   %s%s" % (label, "" if detail is None else "  ← %s" % (detail,)))
    else:
        FAIL += 1
        print("  FAIL %s%s" % (label, "" if detail is None else "  ← %s" % (detail,)))


# ── بدنه‌ای که رابط می‌فرستد (همان قراردادِ toExportRows) ─────────────────────
HEADERS = ["نماد", "حجم (سهم)", "ارزش بازار (ریال)", "سهم حقیقی از ارزش خرید (٪)",
           "سرانه خرید (سهم به ازای هر نفر)", "قیمت صف اول فروش (ریال)",
           "درصد از نشستِ خودِ تابلو؟", "ساعتِ خودِ ردیف (HH:MM)", "منبع داده"]
R1 = ["فولاد", 1479239170, 8784900000000000, 89.6, 715044.0, 0, "نیست", "12:45",
      "TSETMC GetMarketWatch (همان فیدِ زندهٔ تابلو)"]
# نمادی که امروز ردیفی درِ تابلو ندارد: هیچ خانه‌ای صفر نمی‌شود
R2 = ["همراه", None, None, None, None, None, None, None, None]
ROWS = [HEADERS, R1, R2]
META = [["منبع داده", "TSETMC GetMarketWatch — همان فیدِ زندهٔ تابلو که جدول می‌خواند"],
        ["نشستِ داده (d_even)", "20261014"],
        ["ساعتِ فید (h_even)", "13:12"],
        ["زمانِ گرفتنِ خروجی", "2026-10-14 13:15"]]


def xlsx_of(resp):
    return openpyxl.load_workbook(io.BytesIO(resp.body))


def main() -> int:
    print("watch_export_v1106 — گاردِ خروجیِ دیدهبان")

    x = _export._export_build("xlsx", "دیدهبان-بازار", ROWS, meta=META)
    ck(x.status_code == 200, "Excel با کد ۲۰۰ ساخته می‌شود", x.media_type)
    wb = xlsx_of(x)
    ck(wb.sheetnames[0] == "export", "صفحۀ اول همان جدول است", wb.sheetnames)
    ws = wb["export"]
    ck(ws.max_row == 3, "سه سطر: سرِ ستون + دو ردیفِ دیدهبان", ws.max_row)
    ck(ws.max_column == len(HEADERS), "هیچ ستونی درِ Excel افتاده نیست",
       "%d در برابرِ %d" % (ws.max_column, len(HEADERS)))

    hdr = [c.value for c in ws[1]]
    ck(hdr == HEADERS, "سرستون‌هایِ فارسی با واحدِ روشن عیناً می‌نشینند", hdr[:2])

    # ۱) عددها عدد می‌مانند
    row2 = [c.value for c in ws[2]]
    ck(isinstance(row2[1], (int, float)) and not isinstance(row2[1], bool),
       "حجم درِ Excel سلولِ عددی است، نه رشته", type(row2[1]).__name__)
    ck(row2[1] == 1479239170, "عددِ حجم دقیقاً همان است (نه گرد، نه رشتهٔ فارسی)", row2[1])
    ck(isinstance(row2[2], (int, float)) and row2[2] == 8784900000000000,
       "ارزش بازار (ریال) عددِ درشتِ سالم است", row2[2])
    ck(isinstance(row2[3], float) and abs(row2[3] - 89.6) < 1e-9, "٪ هم عدد است", row2[3])

    # ۲) صفرِ واقعی و بی‌داده قاطی نمی‌شوند
    ck(row2[5] == 0, "صفورِ واقعیِ تابلو (صفِ خالی) صفر می‌ماند", row2[5])
    row3 = [c.value for c in ws[3]]
    empties = [i for i, v in enumerate(row3) if v in ("", None)]
    ck(len(empties) >= 7, "ردیفِ بی‌فید با خانۀ تهی می‌آید، نه صفر",
       "%d خانۀ تهی از %d" % (len(empties), len(HEADERS)))
    ck(row3[1] not in (0, "0"), "بی‌داده به صفر تبدیل نشده", repr(row3[1]))
    ck(row3[0] == "همراه", "نمادِ بی‌ردیف ازِ خروجی حذف نشده", row3[0])

    # ۴) ردپا درِ صفحۀ دوم
    ck("منبع و تازگی" in wb.sheetnames, "صفحۀ ردپا کنارِ جدول ساخته شد", wb.sheetnames)
    ms = wb["منبع و تازگی"]
    pairs = {r[0]: r[1] for r in ms.iter_rows(min_row=2, values_only=True)}
    ck(pairs.get("نشستِ داده (d_even)") == "20261014", "زمانِ snapshot درِ فایل ثبت شد",
       pairs.get("نشستِ داده (d_even)"))
    ck("TSETMC" in str(pairs.get("منبع داده") or ""), "منبعِ داده درِ فایل نوشته شد")
    ck(ms["A1"].value == "کلید", "ردپا با سطرِ سرِ خودش می‌آید", ms["A1"].value)

    # ۶) نامِ فایل دوزبانه
    cd = x.headers.get("content-disposition", "")
    ck("filename*=" in cd and "UTF-8''" in cd, "نامِ فارسی با RFC 5987 می‌آید", cd[:60])
    ck('filename="' in cd, "معادلِ ASCII هم هست (سازگاریِ مصرف‌کنندۀ قدیمی)")

    # ۵) CSV
    c = _export._export_build("csv", "دیدهبان-بازار", ROWS, meta=META)
    raw = c.body.decode("utf-8")
    ck(c.body.startswith(b"\xef\xbb\xbf"), "CSV با BOMِ UTF-8 شروع می‌شود", c.media_type)
    text = raw.lstrip("﻿")
    lines = text.rstrip("\r\n").split("\r\n")
    data = [ln for ln in lines if ln and not ln.startswith("#")]
    meta_lines = [ln for ln in lines if ln.startswith("#")]
    ck(len(data) == 3, "ردپا شمارِ ردیف‌هایِ داده را عوض نمی‌کند", len(data))
    ck(len(meta_lines) == len(META), "همۀ ردپا درِ انتهای فایل است", len(meta_lines))
    parsed = list(csv.reader(io.StringIO("\r\n".join(data))))
    ck(parsed[0] == HEADERS, "سرستون‌ها ازِ CSV هم خوانده می‌شوند", parsed[0][:1])
    ck(parsed[2][1] == "", "خانهٔ بی‌داده درِ CSV تهی است، نه صفر", repr(parsed[2][1]))
    ck(any("فولاد" in ln for ln in lines), "واژۀ فارسی درِ فایل می‌ماند")

    # ۷) بی‌meta = رفتارِ پیشین
    x0 = _export._export_build("xlsx", "market", [HEADERS, R1])
    wb0 = xlsx_of(x0)
    ck(wb0.sheetnames == ["export"], "بی‌ردپا، همان یک صفحۀ پیشین برمی‌گردد", wb0.sheetnames)
    ck(wb0["export"].max_row == 2, "و سطرِ داده هم کم نمی‌شود", wb0["export"].max_row)

    # ۸) دروازۀ POST با کنترلِ ورودی
    import fastapi
    try:
        _export.export_data_post({"fmt": "csv", "title": "t", "rows": {"نه": "فهرست"}})
        ck(False, "rowsِ غیرفهرستی باید ۴۲۲ بگیرد")
    except fastapi.HTTPException as e:
        ck(e.status_code == 422, "rowsِ غیرفهرستی ⇒ ۴۲۲", e.status_code)
    try:
        _export.export_data_post({"fmt": "csv", "title": "t",
                                  "rows": [[1]] * (_export.EXPORT_MAX_ROWS + 1)})
        ck(False, "بدنۀ بیش از سقف باید ۴۱۳ بگیرد")
    except fastapi.HTTPException as e:
        ck(e.status_code == 413, "سقفِ ردیف ⇒ ۴۱۳", e.status_code)
    ok = _export.export_data_post({"fmt": "xlsx", "title": "دیدهبان-بازار",
                                   "rows": ROWS, "meta": META})
    ck(getattr(ok, "status_code", 0) == 200 and "منبع و تازگی" in xlsx_of(ok).sheetnames,
       "POSTِ واقعیِ رابط همان فایلِ ردپادار را می‌سازد")

    # ۹) هیچ رازی درِ بدنه نیست
    low = (raw + str(x.body[:200])).lower()
    ck(not any(k in low for k in ("authorization", "set-cookie", "password", "api_key")),
       "هیچ سرِ احرازِ نشستی درِ فایل نیست")

    # کنترلِ منفی: اگر روزی کسی `None` را صفر کند، همین سنسور باید قرمز شود —
    # یعنی سنسور «تهی» را می‌بیند و با صفرِ جعلی یکی نمی‌گوید.
    xp = _export._export_build("xlsx", "t", [HEADERS, [0 if v is None else v for v in R2]])
    vp = [c.value for c in xlsx_of(xp)["export"][2]]
    ck(vp[1] == 0, "کنترلِ سنسور: صفرِ جعلی درِ همان خانۀ تهی شمرده می‌شود", repr(vp[1]))
    zp = _export._export_build("xlsx", "t", [HEADERS, R2])
    vz = [c.value for c in xlsx_of(zp)["export"][2]]
    ck(vz[1] in (None, ""), "و همان خانۀ تهی درِ حالتِ درست، صفر نیست", repr(vz[1]))

    print("\n%d passed / %d failed" % (PASS, FAIL))
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
